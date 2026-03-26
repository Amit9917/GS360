#!/usr/bin/env python3
"""Generate a professionally styled SLA docx document."""

import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

doc = Document()

# -- Page setup --
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# -- Style definitions --
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x2C, 0x5F, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF2, 0xF5, 0xF9)
TABLE_HEADER_BG = "1B2A4A"
TABLE_ALT_BG = "F2F5F9"
BORDER_COLOR = "CCCCCC"

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(10.5)
font.color.rgb = DARK_GRAY
pf = style.paragraph_format
pf.space_after = Pt(6)
pf.space_before = Pt(0)
pf.line_spacing = 1.15

# Heading 1
h1 = doc.styles['Heading 1']
h1.font.name = 'Calibri'
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = NAVY
h1.paragraph_format.space_before = Pt(24)
h1.paragraph_format.space_after = Pt(8)
h1.paragraph_format.keep_with_next = True

# Heading 2
h2 = doc.styles['Heading 2']
h2.font.name = 'Calibri'
h2.font.size = Pt(12)
h2.font.bold = True
h2.font.color.rgb = ACCENT
h2.paragraph_format.space_before = Pt(16)
h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.keep_with_next = True

# -- Helper functions --
def add_para(text, bold=False, italic=False, size=None, color=None, align=None, space_after=None, space_before=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    if bold: run.bold = True
    if italic: run.italic = True
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if align: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None: p.paragraph_format.space_before = Pt(space_before)
    return p

def add_mixed_para(parts, align=None, space_after=None, space_before=None):
    """parts = list of (text, bold, italic, size, color) tuples"""
    p = doc.add_paragraph()
    for text, bold, italic, size, color in parts:
        run = p.add_run(text)
        run.font.name = 'Calibri'
        if bold: run.bold = True
        if italic: run.italic = True
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = color
    if align: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None: p.paragraph_format.space_before = Pt(space_before)
    return p

def add_bullet(text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_GRAY
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.color.rgb = DARK_GRAY
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(1)
    return p

def add_section_heading(number, title):
    p = doc.add_heading(level=1)
    run = p.add_run(f"{number}. {title}")
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    # Add bottom border
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="2C5F8A"/></w:pBdr>')
    pPr.append(pBdr)
    return p

def add_sub_heading(text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    return p

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, val in kwargs.items():
        element = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="single" w:sz="4" w:space="0" w:color="{val}"/>')
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.color.rgb = WHITE
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_shading(cell, TABLE_HEADER_BG)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

    # Data rows
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = DARK_GRAY
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if r % 2 == 1:
                set_cell_shading(cell, TABLE_ALT_BG)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)

    return table

def add_thin_line():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="2" w:space="1" w:color="{BORDER_COLOR}"/></w:pBdr>')
    pPr.append(pBdr)

def add_page_break():
    doc.add_page_break()

def add_note_box(text):
    """Add a highlighted note paragraph"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(9.5)
    run.font.color.rgb = MED_GRAY
    run.italic = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


# ============================================================
# COVER PAGE
# ============================================================
for _ in range(6):
    doc.add_paragraph()

add_para("SERVICE LEVEL AGREEMENT", bold=True, size=28, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_thin_line()
add_para("RIS Software Maintenance & Support Services", bold=False, size=14, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40, space_before=8)

add_para("Submitted to", italic=True, size=10, color=MED_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("Mr. Gagandeep Gupta", bold=True, size=12, color=DARK_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("Scan4Health", size=11, color=DARK_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

add_para("Prepared by", italic=True, size=10, color=MED_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("Jina Code Systems", bold=True, size=12, color=DARK_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("855, Spaze iTech Park, Sector 49, Gurugram", size=10, color=MED_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("March 2026", size=10, color=MED_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

add_page_break()

# ============================================================
# PREAMBLE
# ============================================================
add_para("SERVICE LEVEL AGREEMENT", bold=True, size=16, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

add_para("This Service Level Agreement (\"Agreement\") is entered into between:", space_after=8)

t = doc.add_table(rows=4, cols=2)
t.style = 'Table Grid'
data = [
    ("Service Provider", "Jina Code Systems (\"Provider\")"),
    ("Client", "Scan4Health (\"Client\")"),
    ("Effective Date", "25/03/2026"),
    ("Term", "1 Year from Effective Date"),
]
for i, (label, value) in enumerate(data):
    cell_l = t.rows[i].cells[0]
    cell_r = t.rows[i].cells[1]
    cell_l.text = ""
    cell_r.text = ""
    p_l = cell_l.paragraphs[0]
    run_l = p_l.add_run(label)
    run_l.bold = True
    run_l.font.name = 'Calibri'
    run_l.font.size = Pt(10)
    run_l.font.color.rgb = NAVY
    p_l.paragraph_format.space_before = Pt(4)
    p_l.paragraph_format.space_after = Pt(4)
    set_cell_shading(cell_l, TABLE_ALT_BG)

    p_r = cell_r.paragraphs[0]
    run_r = p_r.add_run(value)
    run_r.font.name = 'Calibri'
    run_r.font.size = Pt(10)
    run_r.font.color.rgb = DARK_GRAY
    p_r.paragraph_format.space_before = Pt(4)
    p_r.paragraph_format.space_after = Pt(4)

t.rows[0].cells[0].width = Inches(2)

add_para("", space_after=4)

# ============================================================
# 1. PURPOSE
# ============================================================
add_section_heading(1, "Purpose")
add_para("This Agreement defines the service levels, support structure, responsibilities, performance commitments, and commercial terms for ongoing maintenance and support of the software system developed and deployed by the Provider for the Client.")
add_para("The objective of this SLA is to ensure uninterrupted system performance, quick issue resolution, and continuous operational efficiency for Scan4Health's diagnostic workflows.")

# ============================================================
# 2. SCOPE OF SERVICES
# ============================================================
add_section_heading(2, "Scope of Services")
add_para("The Provider agrees to deliver comprehensive post-deployment support services, including but not limited to:")

add_sub_heading("2.1 Maintenance & Support")
add_bullet("Preventive maintenance to avoid system failures")
add_bullet("Fixing bugs, defects, and system errors")

add_sub_heading("2.2 Issue Resolution")
add_bullet("Diagnosis and resolution of functional and technical issues")
add_bullet("Database-level troubleshooting (if applicable)")
add_bullet("API and integration-level debugging (existing integrations only)")

add_sub_heading("2.3 System Health Monitoring")
add_bullet("Periodic system health checks (monthly)")
add_bullet("Performance tuning and optimization recommendations")

add_sub_heading("2.4 Technical Assistance")
add_bullet("Support to front-office and operational teams")
add_bullet("Guidance on system usage and best practices")

add_sub_heading("2.5 Minor Enhancements (Fair Usage)")
add_para("Minor enhancements and non-critical improvements may be accommodated under this SLA, subject to the following limits:")
add_bullet("Each individual task must not exceed 4 hours of effort")
add_bullet("A maximum of 20 hours per month of cumulative enhancement effort is included")
add_bullet("Any task exceeding 4 hours, or effort exceeding the monthly cap, will be treated as a Change Request (CR) per Section 12")
add_bullet("Minor enhancements exclude: database schema changes, new workflow creation, new integrations, new module development, and migration tasks")

add_sub_heading("2.6 Important Clarifications")
add_bullet("Any new feature development, module addition, or major workflow change is not included in the base SLA scope and will be treated as a Change Request (CR) as defined in Section 12")
add_bullet("Provider\u2019s support responsibility is limited to the application layer only (software, application logic, UI, and application-level configurations)")
add_bullet("Infrastructure-level issues including but not limited to OS, networking, hardware, and server administration are outside Provider\u2019s scope")

# ============================================================
# 3. PRE-SLA SYSTEM AUDIT
# ============================================================
add_section_heading(3, "Pre-SLA System Audit & Baseline Acceptance")

add_sub_heading("3.1 System Audit")
add_para("Prior to SLA commencement, the Provider shall conduct a system audit (\u201cBaseline Audit\u201d) of the application in its current deployed state. The audit will document:")
add_bullet("Current system version and configuration")
add_bullet("List of known issues, bugs, or defects existing at the time of audit (\u201cKnown Issues\u201d)")
add_bullet("Infrastructure environment and recommended minimum specifications")

add_sub_heading("3.2 Known Issues Register")
add_para("A Known Issues Register will be jointly signed by both parties. Issues listed in this register:")
add_bullet("Will be addressed on a best-effort basis and are not subject to the resolution timelines defined in Section 6")
add_bullet("Major remediation of known issues may be treated as a Change Request (CR) at the Provider\u2019s discretion")

add_sub_heading("3.3 Recommended Infrastructure")
add_para("The Provider will provide a Recommended Infrastructure Specification document. The SLA commitments in this Agreement are contingent upon the Client maintaining infrastructure that meets or exceeds these specifications.")

# ============================================================
# 4. SYSTEM AVAILABILITY
# ============================================================
add_section_heading(4, "System Availability & Infrastructure Dependency")
add_para("The application is deployed on the Client\u2019s in-house infrastructure.")
add_para("No uptime percentage or availability guarantee is provided under this SLA, as system availability is dependent on the Client\u2019s infrastructure which is outside the Provider\u2019s control.", italic=True, size=10, color=MED_GRAY)
add_para("The Provider shall use commercially reasonable efforts to ensure that the software application remains stable and functional under normal operating conditions, provided the Client\u2019s infrastructure meets the Recommended Infrastructure Specification (Section 3.3).")
add_para("System availability is dependent on the Client\u2019s infrastructure, including server uptime, network connectivity, power supply, and third-party services.")

add_sub_heading("Provider shall not be responsible for downtime caused by:")
add_bullet("Server failures or hardware issues")
add_bullet("Network outages")
add_bullet("Improper or sub-standard infrastructure configuration")
add_bullet("Third-party software, OS, or service failures")
add_bullet("Client-side modifications to the system (see Section 8)")

add_para("In case of infrastructure-related downtime, the Provider will assist in diagnosis and support resolution on a best-effort basis, but such assistance will not count toward SLA resolution timelines.")

# ============================================================
# 5. SUPPORT HOURS
# ============================================================
add_section_heading(5, "Support Hours")

add_sub_heading("5.1 Standard Support")
add_para("Monday to Friday, 09:30 AM to 5:30 PM (IST)")

add_sub_heading("5.2 Emergency Support")
add_para("24/7 support available only for Severity 1 (Critical) issues. Emergency support outside standard business hours may be charged at \u20b93,000/hour if the issue is subsequently determined to be caused by Client infrastructure, unauthorized changes, or factors outside Provider\u2019s application scope.")

add_sub_heading("5.3 Public Holidays")
add_bullet("Indian national gazetted holidays and Haryana state holidays are excluded from business hours")
add_bullet("Response and resolution timelines that reference \u201cbusiness hours\u201d or \u201cworking days\u201d do not include public holidays")
add_bullet("A holiday calendar will be shared with the Client at the start of each calendar year")

# ============================================================
# 6. ISSUE CLASSIFICATION
# ============================================================
add_section_heading(6, "Issue Classification, Response & Resolution")

add_sub_heading("6.1 Severity Levels")
add_styled_table(
    ["Severity", "Description", "Target Response", "Target Resolution"],
    [
        ("Severity 1 \u2014 Critical", "Complete system downtime or core operations blocked", "Within 2 hours", "Within 12 hours"),
        ("Severity 2 \u2014 High", "Major functionality impacted with no workaround", "Within 8 business hours", "Within 2 working days"),
        ("Severity 3 \u2014 Medium", "Partial functionality issue, workaround available", "Within 1 working day", "3 to 5 working days"),
        ("Severity 4 \u2014 Low", "Cosmetic or UI issues, non-critical improvements", "Within 2 working days", "Next planned release"),
    ],
    col_widths=[1.5, 2.2, 1.4, 1.4]
)

add_note_box("All response and resolution times are target / best-effort timelines, not guaranteed commitments. No financial penalties shall be imposed on the Provider for failure to meet target resolution times. Resolution clocks run only during business hours unless Severity 1 emergency support is engaged.")

add_sub_heading("6.2 Severity Classification")
add_bullet("Final severity classification shall be mutually agreed between the Provider and the Client at the time of issue logging")
add_bullet("In case of disagreement, the Provider\u2019s assessment shall prevail until reviewed at the next escalation level")
add_bullet("The Provider reserves the right to reclassify severity upon investigation")

add_sub_heading("6.3 SLA Clock Suspension")
add_para("The response and resolution timers shall be automatically paused when:")
add_bullet("The Provider is awaiting information, access, credentials, or approvals from the Client")
add_bullet("The issue requires action from a third-party vendor or service outside Provider\u2019s control")
add_bullet("The Client\u2019s infrastructure is unavailable or inaccessible")
add_bullet("The Client requests a hold or postponement of work on a ticket")
add_para("The clock resumes upon receipt of the required information or access.")

add_sub_heading("6.4 Workaround vs Permanent Fix")
add_bullet("Resolution may include a temporary workaround that restores functionality, even if a permanent fix is pending")
add_bullet("A workaround that restores business operations is a valid and complete resolution for SLA timeline purposes")
add_bullet("Permanent fixes will be scheduled in the next planned release at the Provider\u2019s discretion")
add_bullet("The Provider is under no obligation to provide a permanent fix within a specific timeframe once a workaround has been delivered")

add_sub_heading("6.5 Fair Usage of Support")
add_bullet("Maximum of 30 support hours per month for issue resolution and maintenance (excluding minor enhancements in Section 2.5)")
add_bullet("If the Client exceeds the monthly cap, additional hours may be charged at \u20b92,000/hour, non-critical tickets may be deferred, or an upgraded plan may be proposed")
add_bullet("Unused hours do not carry over to subsequent months")

add_sub_heading("6.6 Right to Decline")
add_para("The Provider reserves the right to decline or defer any support request that falls outside scope, relates to unauthorized changes or infrastructure failures, or would require effort disproportionate to the SLA fee. In such cases, the Provider may offer to address the request as a Change Request.")

# ============================================================
# 7. SUPPORT CHANNELS
# ============================================================
add_section_heading(7, "Support Channels")
add_bullet("Email Support: tech@attention.sh")
add_bullet("Ticketing System (if applicable)")
add_para("Only issues logged through official support channels will be tracked under this SLA. Issues communicated via informal channels will not be considered valid tickets unless subsequently logged officially. Each ticket must include: issue description, severity assessment, steps to reproduce, and relevant screenshots or logs.")

# ============================================================
# 8. CLIENT RESPONSIBILITIES
# ============================================================
add_section_heading(8, "Client Responsibilities")
add_para("The Client agrees to:")
add_bullet("Provide detailed issue descriptions with screenshots and logs where applicable")
add_bullet("Ensure timely access to systems, servers, or credentials (within 2 business hours for Severity 1; within 1 business day for others)")
add_bullet("Nominate a single point of contact (SPOC) for coordination")
add_bullet("Not make unauthorized changes to codebase, database, or infrastructure \u2014 unauthorized changes will void Provider\u2019s obligation to resolve resulting issues")
add_bullet("Maintain infrastructure in accordance with the Recommended Infrastructure Specification (Section 3.3)")
add_bullet("Ensure adequate data backup procedures are in place (backups are the Client\u2019s sole responsibility)")

# ============================================================
# 9. ESCALATION MATRIX
# ============================================================
add_section_heading(9, "Escalation Matrix")
add_styled_table(
    ["Level", "Role", "Target Response Time"],
    [
        ("Level 1", "Support Engineer", "As per Section 6.1"),
        ("Level 2", "Project Manager", "Within 4 business hours of escalation"),
        ("Level 3", "Senior Management", "Within 1 business day of escalation"),
    ],
    col_widths=[1.2, 2.5, 2.8]
)
add_para("Escalation to the next level may be initiated by the Client if the current level has not acknowledged or responded within the above timelines.", space_before=6)

# ============================================================
# 10. EXCLUSIONS
# ============================================================
add_section_heading(10, "Exclusions")
add_para("The following are explicitly excluded from this SLA:")
add_bullet("Development of new modules or features")
add_bullet("Major workflow redesign")
add_bullet("New third-party integrations")
add_bullet("Hardware or infrastructure failures")
add_bullet("Issues caused due to misuse, unauthorized changes, or tampering by Client or third parties")
add_bullet("External system or API failures beyond Provider\u2019s control")
add_bullet("Database administration tasks (backups, replication, server tuning)")
add_bullet("Data migration or bulk data operations")
add_bullet("Issues arising from Client\u2019s failure to maintain Recommended Infrastructure Specifications")
add_bullet("Force majeure events")
add_bullet("Any issue the Provider reasonably determines to be outside the application layer")

# ============================================================
# 11. DATA SECURITY
# ============================================================
add_section_heading(11, "Data Security & Confidentiality")

add_sub_heading("11.1 Confidentiality")
add_bullet("All Client data will be treated as strictly confidential")
add_bullet("No data sharing with third parties without written consent")
add_bullet("Both parties agree to maintain confidentiality of proprietary and business information")

add_sub_heading("11.2 Access Control")
add_bullet("Provider personnel will access Client systems only through authorized credentials provided by the Client")
add_bullet("Access will be limited to personnel directly involved in support and maintenance")
add_bullet("Access will be through encrypted remote connections (SSH/VPN) only")

add_sub_heading("11.3 Data Handling")
add_bullet("The Provider will not copy, export, or store Client data externally unless required for troubleshooting with prior written approval")
add_bullet("Any temporarily copied data will be securely deleted within a reasonable timeframe after issue resolution")

add_sub_heading("11.4 Security Practices")
add_para("The Provider will follow industry-standard security practices including encrypted remote access, secure credential handling, principle of least privilege, and multi-factor authentication where supported.")
add_para("The Provider is not responsible for data backups, disaster recovery, data loss caused by infrastructure failures or Client actions, or security breaches resulting from Client\u2019s failure to maintain secure infrastructure.")

# ============================================================
# 12. CHANGE REQUESTS
# ============================================================
add_section_heading(12, "Change Requests & New Features")
add_para("Any requirement outside the defined SLA scope \u2014 including new features, workflow changes, additional reports, new integrations, database schema changes, or enhancements exceeding Section 2.5 limits \u2014 will be treated as a Change Request (CR).")

add_sub_heading("12.1 Commercial Terms")
add_styled_table(
    ["Item", "Detail"],
    [
        ("CR Hourly Rate", "\u20b92,000/hour (exclusive of applicable taxes)"),
        ("Approval", "Written approval required before work begins"),
        ("Estimate Variance", "Up to 20% overrun without re-approval"),
        ("Invoice Terms", "Payable within 15 days of invoice date"),
    ],
    col_widths=[2.0, 4.5]
)

add_sub_heading("12.2 Delivery Timelines")
add_bullet("Timelines depend on scope and complexity, mutually agreed before development begins")
add_bullet("CR work will not affect or delay ongoing SLA support obligations")

# ============================================================
# 13. PAYMENT TERMS
# ============================================================
add_section_heading(13, "Payment Terms")

add_sub_heading("13.1 SLA Fee")
add_styled_table(
    ["Item", "Amount"],
    [
        ("Annual SLA Fee (excl. GST)", "\u20b91,50,000/-"),
        ("GST (18%)", "\u20b927,000/-"),
        ("Total Payable", "\u20b91,77,000/-"),
        ("Payment Terms", "100% advance before SLA commencement"),
    ],
    col_widths=[3.0, 3.5]
)
add_para("SLA validity begins only after payment is received and confirmed. All amounts are exclusive of applicable taxes unless stated otherwise.", space_before=6)

add_sub_heading("13.2 Delay in Payment")
add_bullet("Written reminder after 7 days past due date")
add_bullet("All support services will be suspended after 15 days past due date")
add_bullet("Issues arising during suspension will not be covered and Provider bears no liability")

add_sub_heading("13.3 Reinstatement")
add_para("Upon receipt of overdue payment, services will be reinstated within 2 business days. The SLA term will not be extended for the suspension period.")

# ============================================================
# 14. INTELLECTUAL PROPERTY
# ============================================================
add_section_heading(14, "Intellectual Property")

add_sub_heading("14.1 Provider IP")
add_para("The Provider retains full ownership of all intellectual property in the core software, including source code, architecture, libraries, frameworks, tools, and methodologies. The Client is granted a non-exclusive, non-transferable license to use the software as deployed for the duration of this Agreement.")

add_sub_heading("14.2 Change Request IP")
add_para("IP created as part of Change Requests shall be owned by the Provider unless explicitly agreed otherwise in writing. Reusable components developed during CR execution remain the exclusive property of the Provider.")

add_sub_heading("14.3 Client Data")
add_para("All data entered, generated, or stored by the Client within the system remains the exclusive property of the Client. The Provider shall have no claim over Client data.")

# ============================================================
# 15. TERMINATION
# ============================================================
add_section_heading(15, "Termination")

add_sub_heading("15.1 Termination by Provider")
add_para("The Provider may terminate this Agreement with 30 days written notice for:")
add_bullet("Non-payment beyond 30 days from due date")
add_bullet("Breach of confidentiality or IP terms by the Client")
add_bullet("Unauthorized modifications that fundamentally compromise the software")
add_bullet("Repeated failure to maintain Recommended Infrastructure Specifications")
add_bullet("Persistent fair usage exceedance where Client declines an upgraded plan")
add_bullet("Any material breach uncured for 15 days after written notice")

add_sub_heading("15.2 Termination by Client")
add_para("The Client may terminate with 30 days written notice. However:")
add_bullet("No refund will be provided for the remaining SLA term")
add_bullet("If terminated within the first 6 months, the full annual SLA fee becomes due")
add_bullet("All outstanding CR invoices become immediately payable upon termination")

add_sub_heading("15.3 Post-Termination")
add_para("Upon termination, the Provider will cease all support services, revoke access to Provider tools, and return or securely delete Client data within 30 days. The Client\u2019s software license is governed by the original license agreement, separate from this SLA.")

# ============================================================
# 16. LIMITATION OF LIABILITY
# ============================================================
add_section_heading(16, "Limitation of Liability")
add_para("Total aggregate liability of the Provider shall not exceed 50% of the annual SLA fee paid in the 12 months preceding the claim.", bold=True)
add_para("The Provider shall not be liable for indirect, incidental, consequential, special, or punitive damages; loss of revenue, profits, data, or business opportunities; business interruption; data loss due to infrastructure failures or Client actions; delays caused by Client\u2019s failure to provide timely access; issues from unauthorized changes; or damages from reliance on Provider\u2019s recommendations.")
add_para("This limitation applies regardless of the form of action. No SLA penalties, service credits, or liquidated damages shall apply. The escalation process (Section 9) is the Client\u2019s sole and exclusive remedy for service level failures.")

# ============================================================
# 17. INDEMNIFICATION
# ============================================================
add_section_heading(17, "Indemnification")
add_para("The Client shall indemnify, defend, and hold harmless the Provider, its officers, employees, and agents from any claims, losses, damages, liabilities, or expenses arising from:")
add_bullet("The Client\u2019s unauthorized modifications to software, database, or infrastructure")
add_bullet("Failure to maintain infrastructure per Recommended Specifications")
add_bullet("Third-party claims related to the Client\u2019s use of the software or data")
add_bullet("The Client\u2019s breach of any terms of this Agreement")
add_para("This indemnification obligation survives termination of this Agreement.", italic=True)

# ============================================================
# 18. RENEWAL
# ============================================================
add_section_heading(18, "Renewal")
add_bullet("Renewable annually by mutual written consent at least 30 days before expiry")
add_bullet("The Provider reserves the right to revise fees, terms, and scope at renewal (30 days advance notice)")
add_bullet("Agreement will lapse if no renewal terms are agreed before expiry")
add_bullet("Renewal contingent upon clearance of all outstanding dues")

# ============================================================
# 19. DISPUTE RESOLUTION
# ============================================================
add_section_heading(19, "Dispute Resolution")
add_bullet("Disputes shall first be resolved through good-faith negotiation")
add_bullet("If unresolved within 30 days, referred to arbitration under the Arbitration and Conciliation Act, 1996")
add_bullet("Seat of arbitration: Gurugram, Haryana. Language: English")
add_bullet("Each party bears its own costs unless the arbitrator orders otherwise")

# ============================================================
# 20. GOVERNING LAW
# ============================================================
add_section_heading(20, "Governing Law")
add_para("This Agreement shall be governed by the laws of India. The courts of Gurugram, Haryana shall have exclusive jurisdiction over any matters not subject to arbitration.")

# ============================================================
# 21. FORCE MAJEURE
# ============================================================
add_section_heading(21, "Force Majeure")
add_bullet("Neither party liable for delays due to natural disasters, pandemics, war, government orders, widespread outages, or events beyond reasonable control")
add_bullet("Affected party shall notify the other within 5 business days")
add_bullet("If force majeure exceeds 60 consecutive days, either party may terminate with 15 days notice, without penalty")
add_bullet("All SLA timelines suspended during force majeure; no refund of SLA fees for the suspension period")

# ============================================================
# 22. NOTICES
# ============================================================
add_section_heading(22, "Notices")
add_para("All formal notices shall be delivered via email to designated contacts and via registered post or courier to registered addresses. A notice is deemed delivered on email delivery confirmation or 3 business days after postal dispatch, whichever is earlier.")
add_para("Designated contacts:")
add_bullet("Provider: Amit Joshi, COO \u2014 _________________________ (email)")
add_bullet("Client: _________________________ \u2014 _________________________ (email)")
add_para("Either party may update contacts with 7 days written notice.")

# ============================================================
# 23. ENTIRE AGREEMENT
# ============================================================
add_section_heading(23, "Entire Agreement & Amendments")
add_bullet("This Agreement and its Annexures constitute the entire agreement, superseding all prior discussions")
add_bullet("Amendments require written consent signed by authorized representatives of both parties")
add_bullet("Failure to enforce a provision is not a waiver of that provision")
add_bullet("If any provision is unenforceable, remaining provisions continue in full force (severability)")

add_page_break()

# ============================================================
# 24. ACCEPTANCE
# ============================================================
add_section_heading(24, "Acceptance")
add_para("By signing below, both parties acknowledge and agree to the terms and conditions outlined in this Service Level Agreement.", space_after=20)

# Provider signature block
add_para("For Jina Code Systems", bold=True, size=12, color=NAVY, space_after=12)
sig_table = doc.add_table(rows=4, cols=2)
sig_data = [("Name", "Amit Joshi"), ("Designation", "COO"), ("Signature", "_________________________"), ("Date", "25/03/2026")]
for i, (label, value) in enumerate(sig_data):
    c1 = sig_table.rows[i].cells[0]
    c2 = sig_table.rows[i].cells[1]
    c1.text = ""
    c2.text = ""
    r1 = c1.paragraphs[0].add_run(label)
    r1.font.name = 'Calibri'; r1.font.size = Pt(10); r1.font.color.rgb = MED_GRAY
    r2 = c2.paragraphs[0].add_run(value)
    r2.font.name = 'Calibri'; r2.font.size = Pt(10); r2.font.color.rgb = DARK_GRAY

add_para("", space_after=16)

# Client signature block
add_para("For Scan4Health", bold=True, size=12, color=NAVY, space_after=12)
sig_table2 = doc.add_table(rows=4, cols=2)
sig_data2 = [("Name", "_________________________"), ("Designation", "_________________________"), ("Signature", "_________________________"), ("Date", "_________________________")]
for i, (label, value) in enumerate(sig_data2):
    c1 = sig_table2.rows[i].cells[0]
    c2 = sig_table2.rows[i].cells[1]
    c1.text = ""
    c2.text = ""
    r1 = c1.paragraphs[0].add_run(label)
    r1.font.name = 'Calibri'; r1.font.size = Pt(10); r1.font.color.rgb = MED_GRAY
    r2 = c2.paragraphs[0].add_run(value)
    r2.font.name = 'Calibri'; r2.font.size = Pt(10); r2.font.color.rgb = DARK_GRAY

add_page_break()

# ============================================================
# ANNEXURE A
# ============================================================
add_section_heading("A", "Known Issues Register")
add_note_box("To be completed during Pre-SLA System Audit (Section 3)")
add_styled_table(
    ["#", "Issue Description", "Severity", "Date Identified", "Remarks"],
    [("1", "", "", "", ""), ("2", "", "", "", ""), ("3", "", "", "", ""), ("4", "", "", "", ""), ("5", "", "", "", "")],
    col_widths=[0.4, 2.5, 1.0, 1.1, 1.5]
)
add_para("", space_after=16)
add_para("Acknowledged by Provider: _________________________     Date: ___________")
add_para("Acknowledged by Client: _________________________     Date: ___________")

add_page_break()

# ============================================================
# ANNEXURE B
# ============================================================
add_section_heading("B", "Recommended Infrastructure Specification")
add_note_box("To be completed during Pre-SLA System Audit (Section 3)")
add_styled_table(
    ["Component", "Minimum Specification"],
    [("Server OS", ""), ("CPU", ""), ("RAM", ""), ("Storage", ""), ("Network", ""), ("Database", ""), ("Other", "")],
    col_widths=[2.0, 4.5]
)

# ============================================================
# SAVE
# ============================================================
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SLA_Final_Professional.docx")
doc.save(output_path)
print(f"Saved to {output_path}")
